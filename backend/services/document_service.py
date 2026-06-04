import base64
import io

import fitz
from docx import Document
from pptx import Presentation


SUPPORTED_EXTS = {"pdf", "docx", "pptx", "png", "jpg", "jpeg"}


def extract_content(file_bytes: bytes, ext: str) -> tuple:
    """Returns (content, metadata). content is str for text-based or list[base64] for vision."""
    ext = ext.lower()
    if ext == "pdf":
        return _extract_pdf(file_bytes)
    if ext == "docx":
        return _extract_docx(file_bytes)
    if ext == "pptx":
        return _extract_pptx(file_bytes)
    if ext in ("png", "jpg", "jpeg"):
        return _encode_image(file_bytes, ext), {"mode": "vision", "pages": 1}
    raise ValueError(f"Formato no soportado: {ext}")


def _extract_pdf(file_bytes: bytes) -> tuple:
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages_text = [page.get_text() for page in doc]
    full_text = "\n\n".join(pages_text).strip()
    metadata = {"pages": len(doc)}

    if len(full_text) < 100:
        images = []
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
            images.append(base64.standard_b64encode(pix.tobytes("png")).decode())
        return images, {**metadata, "mode": "vision", "is_scanned": True}

    return full_text, {**metadata, "mode": "text", "is_scanned": False}


def _extract_docx(file_bytes: bytes) -> tuple:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    metadata = {"tables": len(doc.tables), "mode": "text", "pages": None}
    return "\n\n".join(paragraphs), metadata


def _extract_pptx(file_bytes: bytes) -> tuple:
    prs = Presentation(io.BytesIO(file_bytes))
    slides = []
    for i, slide in enumerate(prs.slides):
        texts = [s.text for s in slide.shapes if hasattr(s, "text") and s.text.strip()]
        if texts:
            slides.append(f"[Slide {i + 1}]\n" + "\n".join(texts))
    metadata = {"slides": len(prs.slides), "mode": "text", "pages": len(prs.slides)}
    return "\n\n".join(slides), metadata


def _encode_image(file_bytes: bytes, ext: str) -> list:
    return [base64.standard_b64encode(file_bytes).decode()]
