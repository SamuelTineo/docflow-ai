import io

from docx import Document
from pptx import Presentation


LANGUAGE_NAMES = {
    "en": "English", "es": "Spanish", "pt": "Portuguese",
    "fr": "French",  "de": "German",  "it": "Italian",
    "zh": "Chinese", "ja": "Japanese",
}

OUTPUT_MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def build_output_doc(original_bytes: bytes, ext: str, translated_text: str, filename: str) -> tuple:
    """Returns (output_bytes, output_ext, mime_type)."""
    ext = ext.lower()
    if ext == "docx":
        return _rebuild_docx(original_bytes, translated_text), "docx", OUTPUT_MIME["docx"]
    if ext == "pptx":
        return _rebuild_pptx(original_bytes, translated_text), "pptx", OUTPUT_MIME["pptx"]
    # PDF / images → create simple DOCX
    return _create_docx(translated_text, filename), "docx", OUTPUT_MIME["docx"]


def _rebuild_docx(original_bytes: bytes, translated_text: str) -> bytes:
    doc = Document(io.BytesIO(original_bytes))
    translated_paras = [p.strip() for p in translated_text.split("\n") if p.strip()]
    orig_paras = [p for p in doc.paragraphs if p.text.strip()]

    for i, para in enumerate(orig_paras):
        if i < len(translated_paras):
            for run in para.runs:
                run.text = ""
            if para.runs:
                para.runs[0].text = translated_paras[i]
            else:
                para.text = translated_paras[i]

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _rebuild_pptx(original_bytes: bytes, translated_text: str) -> bytes:
    prs = Presentation(io.BytesIO(original_bytes))
    translated_blocks = [b.strip() for b in translated_text.split("\n\n") if b.strip()]
    block_idx = 0

    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text.strip():
                if block_idx < len(translated_blocks):
                    tf = shape.text_frame
                    lines = translated_blocks[block_idx].split("\n")
                    for i, para in enumerate(tf.paragraphs):
                        for run in para.runs:
                            run.text = ""
                        if i < len(lines) and para.runs:
                            para.runs[0].text = lines[i]
                    block_idx += 1

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _create_docx(text: str, filename: str) -> bytes:
    doc = Document()
    doc.add_heading(f"Translated: {filename}", level=1)
    for para in text.split("\n\n"):
        para = para.strip()
        if para:
            doc.add_paragraph(para)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
